# 提取 PPTX 全部文字内容：每页的文本、表格、图片占位、演讲者备注。
# 用法: python pptx_text.py <src.pptx> <out.txt>
# 依赖: pip install python-pptx
import sys, io
from pptx import Presentation

out = io.open(sys.argv[2], "w", encoding="utf-8")
prs = Presentation(sys.argv[1])
out.write("slide size: %s x %s\n" % (prs.slide_width, prs.slide_height))
for i, slide in enumerate(prs.slides, 1):
    out.write("\n===== SLIDE %d =====\n" % i)
    for sh in slide.shapes:
        if sh.has_text_frame:
            txt = sh.text_frame.text.strip()
            if txt:
                out.write("[TEXT] %s\n" % txt.replace("\n", " \u23ce "))
        if sh.shape_type == 13:   # PICTURE
            out.write("[PICTURE] name=%s pos=(%s,%s) size=(%s,%s)\n"
                      % (sh.name, sh.left, sh.top, sh.width, sh.height))
        if sh.has_table:
            out.write("[TABLE]\n")
            for r in sh.table.rows:
                out.write("  | " + " | ".join(c.text.strip() for c in r.cells) + " |\n")
    try:
        if slide.has_notes_slide:
            ntf = slide.notes_slide.notes_text_frame
            if ntf is not None:                      # notes_text_frame 可能为 None
                nt = ntf.text.strip()
                if nt:
                    out.write("[NOTES] %s\n" % nt)
    except Exception:
        pass
out.close()
print("done ->", sys.argv[2])
