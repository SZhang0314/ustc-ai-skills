# 从幻灯片 PDF 截取"完整幻灯片"作为讲义插图：
#   1) 整页截图（保证任何图形/文字不被截断——不要按形状 bbox 或固定比例裁！）；
#   2) 按 PPTX 中页码/页脚形状的实际 EMU 坐标，把右下角页码涂白；
#   3) 裁四周白边（trim 不触碰墨迹）。
# 用法: python make_figs.py <figs.py配置文件>
#   配置文件（figs.py）需定义：
#     FIGS = { 输出名: (幻灯片png文件名, ...可选), ... }
#     SLIDES = 幻灯片png目录;  OUT = 输出目录
#     PPTXS = { ppt序号: pptx路径 };  SLIDE_EMU_W/H = 幻灯片EMU尺寸
# 依赖: pip install python-pptx pillow numpy
import os
from pptx import Presentation
from PIL import Image, ImageDraw
import numpy as np

def footer_boxes_emu(slide):
    """右下角页码 / 通栏页脚区域（EMU 坐标）"""
    boxes = []
    for sh in slide.shapes:
        try:
            l, t, w, h = sh.left, sh.top, sh.width, sh.height
        except Exception:
            continue
        if None in (l, t, w, h):
            continue
        l, t, r, b = int(l), int(t), int(l + w), int(t + h)
        if l > 11000000 and t > 6200000:            # 右下角页码（16:9 幻灯片经验值）
            boxes.append((l, t, r, b))
        elif t > 6600000 and (r - l) > 10000000:    # 通栏底条（封面常见）
            boxes.append((l, t, r, b))
    return boxes

def paint_white(img, boxes_emu, sc, pad_px=6):
    d = ImageDraw.Draw(img)
    for (l, t, r, b) in boxes_emu:
        x0 = max(0, int(l * sc) - pad_px); y0 = max(0, int(t * sc) - pad_px)
        x1 = min(img.width, int(r * sc) + pad_px); y1 = min(img.height, int(b * sc) + pad_px)
        d.rectangle([x0, y0, x1, y1], fill=(255, 255, 255))
    return img

def trim_white(img, tol=248, pad=10):
    a = np.asarray(img.convert("L"))
    mask = a < tol
    ys = np.where(mask.any(axis=1))[0]
    xs = np.where(mask.any(axis=0))[0]
    if len(ys) == 0 or len(xs) == 0:
        return img
    x0 = max(0, int(xs[0]) - pad); y0 = max(0, int(ys[0]) - pad)
    x1 = min(img.width, int(xs[-1]) + pad); y1 = min(img.height, int(ys[-1]) + pad)
    return img.crop((x0, y0, x1, y1))

if __name__ == "__main__":
    import importlib.util
    spec = importlib.util.spec_from_file_location("figs_cfg", sys.argv[1])
    cfg = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cfg)

    _cache = {}
    def get_slides(pn):
        if pn not in _cache:
            _cache[pn] = list(Presentation(cfg.PPTXS[pn]).slides)
        return _cache[pn]

    os.makedirs(cfg.OUT, exist_ok=True)
    for name, entry in sorted(cfg.FIGS.items()):
        fn = entry[0] if isinstance(entry, tuple) else entry
        pn, sno = entry[1], entry[2] if isinstance(entry, tuple) and len(entry) > 2 else (None, None)
        png = os.path.join(cfg.SLIDES, fn)
        img = Image.open(png).convert("RGB")
        W, H = img.size
        sc = W / cfg.SLIDE_EMU_W
        if pn is not None:
            fb = footer_boxes_emu(get_slides(pn)[sno - 1])
            img = paint_white(img, fb, sc)
        img = trim_white(img)
        img.save(os.path.join(cfg.OUT, name + ".png"))
        print("%-24s <- %-14s %5dx%-5d footer_painted=%s"
              % (name, fn, img.width, img.height, len(fb) if pn else "n/a"))
